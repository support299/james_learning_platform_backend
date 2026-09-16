"""Background jobs for pptx import: as one slideshow lesson's slides, or as
one image lesson per slide with lesson-to-lesson jump hotspots.

Both jobs carry over the source deck's real "Go to page or object" click
interactions (PowerPoint/Impress's native slide-jump action) as hotspots when
present. Confirmed against a hand-authored deck: python-pptx's
`shape.click_action.target_slide` reads exactly that interaction type, but a
LibreOffice-exported file can have a shape whose click action points at a
relationship id missing from that slide's .rels part — read defensively,
per-shape, so one bad shape doesn't sink the whole import (see
`_extract_click_targets`).
"""

import io
import logging
import subprocess
import tempfile
from pathlib import Path

from celery import shared_task
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from pdf2image import convert_from_path
from pptx import Presentation

from .models import Course, Lesson, SlideshowSlide
from .slugs import unique_slug

logger = logging.getLogger(__name__)


def _pptx_to_page_images(pptx_path):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [
                'soffice', '--headless', '--convert-to', 'pdf',
                '--outdir', tmp, pptx_path,
            ],
            check=True, timeout=120, capture_output=True,
        )
        pdf_path = Path(tmp) / (Path(pptx_path).stem + '.pdf')
        pages = convert_from_path(str(pdf_path), dpi=150)
    return Presentation(pptx_path), pages


def _extract_click_targets(prs, log_context):
    """Per-slide list of draft hotspots, keyed by pptx slide order — each
    `target` is still a pptx slide index here, translated to a real row
    id/slug by the caller once the destination rows exist."""
    sw, sh = prs.slide_width, prs.slide_height
    draft_hotspots = []
    skipped = 0
    for slide in prs.slides:
        spots = []
        for shape in slide.shapes:
            try:
                target = shape.click_action.target_slide
            except KeyError:
                skipped += 1
                continue
            if target is None:
                continue
            spots.append({
                'x': shape.left / sw,
                'y': shape.top / sh,
                'w': shape.width / sw,
                'h': shape.height / sh,
                'target': prs.slides.index(target),
            })
        draft_hotspots.append(spots)
    if skipped:
        logger.warning(
            '%s: skipped %d shape(s) with a broken click-action relationship',
            log_context, skipped,
        )
    return draft_hotspots


@shared_task(name='courses.import_slideshow_pptx', bind=True, max_retries=0, acks_late=True)
def import_slideshow_pptx(self, lesson_id, tmp_relpath):
    """No retry: a bad pptx or a soffice crash is deterministic, so retrying
    just wastes time — the fix is a different file, not another attempt.
    acks_late so a worker dying mid-run (as happened once in production)
    gets the task redelivered to another worker instead of losing it —
    safe since the whole pipeline is idempotent (delete-and-recreate)."""
    try:
        lesson = Lesson.objects.get(pk=lesson_id)
    except Lesson.DoesNotExist:
        logger.warning('Slideshow import for lesson %s: lesson no longer exists', lesson_id)
        default_storage.delete(tmp_relpath)
        return
    pptx_path = default_storage.path(tmp_relpath)

    try:
        prs, pages = _pptx_to_page_images(pptx_path)
        draft_hotspots = _extract_click_targets(
            prs, f'Slideshow import for lesson {lesson_id}'
        )
        if not Lesson.objects.filter(pk=lesson_id).exists():
            logger.warning(
                'Slideshow import for lesson %s: lesson was deleted mid-import, aborting', lesson_id,
            )
            return

        lesson.slideshow_slides.all().delete()
        created = []
        for i, page in enumerate(pages):
            buf_name = f'slide_{i}.png'
            slide_obj = SlideshowSlide(lesson=lesson, order=i, hotspots=[])
            buf = io.BytesIO()
            page.save(buf, format='PNG')
            slide_obj.image.save(buf_name, ContentFile(buf.getvalue()), save=False)
            slide_obj.save()
            created.append(slide_obj)

        order_to_id = {s.order: s.id for s in created}
        for slide_obj, spots in zip(created, draft_hotspots):
            resolved = [
                {**spot, 'target': order_to_id[spot['target']]}
                for spot in spots
                if spot['target'] in order_to_id
            ]
            if resolved:
                slide_obj.hotspots = resolved
                slide_obj.save(update_fields=['hotspots'])

        lesson.import_status = Lesson.ImportStatus.DONE
        lesson.import_error = ''
        lesson.save(update_fields=['import_status', 'import_error'])
    except Exception:
        logger.exception('Slideshow import failed for lesson %s', lesson_id)
        lesson.import_status = Lesson.ImportStatus.FAILED
        lesson.import_error = (
            'Could not process this file — check it is a valid, non-corrupted '
            '.pptx and try again. Full details are in the server logs.'
        )
        lesson.save(update_fields=['import_status', 'import_error'])
    finally:
        default_storage.delete(tmp_relpath)


def _slide_title(slide, index):
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        text = slide.shapes.title.text_frame.text.strip()
        if text:
            return text
    return f'Slide {index + 1}'


@shared_task(name='courses.import_course_lessons_pptx', bind=True, max_retries=0, acks_late=True)
def import_course_lessons_pptx(self, course_id, tmp_relpath):
    try:
        course = Course.objects.get(pk=course_id)
    except Course.DoesNotExist:
        logger.warning('Course lesson import for course %s: course no longer exists', course_id)
        default_storage.delete(tmp_relpath)
        return
    pptx_path = default_storage.path(tmp_relpath)

    try:
        prs, pages = _pptx_to_page_images(pptx_path)
        draft_hotspots = _extract_click_targets(
            prs, f'Course lesson import for course {course_id}'
        )
        if not Course.objects.filter(pk=course_id).exists():
            logger.warning(
                'Course lesson import for course %s: course was deleted mid-import, aborting', course_id,
            )
            return

        with transaction.atomic():
            existing_slugs = set(course.lessons.values_list('slug', flat=True))
            next_order = course.lessons.count()
            created = []
            for i, (slide, page) in enumerate(zip(prs.slides, pages)):
                title = _slide_title(slide, i)
                slug = unique_slug(title, existing_slugs)
                existing_slugs.add(slug)

                lesson_obj = Lesson(
                    course=course, slug=slug, title=title,
                    lesson_type=Lesson.Type.IMAGE, order=next_order + i,
                )
                buf = io.BytesIO()
                page.save(buf, format='PNG')
                lesson_obj.image.save(f'slide_{i}.png', ContentFile(buf.getvalue()), save=False)
                lesson_obj.save()
                created.append(lesson_obj)

            order_to_slug = {i: lesson_obj.slug for i, lesson_obj in enumerate(created)}
            for lesson_obj, spots in zip(created, draft_hotspots):
                resolved = [
                    {**spot, 'target': order_to_slug[spot['target']]}
                    for spot in spots
                    if spot['target'] in order_to_slug
                ]
                if resolved:
                    lesson_obj.hotspots = resolved
                    lesson_obj.save(update_fields=['hotspots'])

            course.save()

        course.import_status = Course.ImportStatus.DONE
        course.import_error = ''
        course.save(update_fields=['import_status', 'import_error'])
    except Exception:
        logger.exception('Course lesson import failed for course %s', course_id)
        course.import_status = Course.ImportStatus.FAILED
        course.import_error = (
            'Could not process this file — check it is a valid, non-corrupted '
            '.pptx and try again. Full details are in the server logs.'
        )
        course.save(update_fields=['import_status', 'import_error'])
    finally:
        default_storage.delete(tmp_relpath)
